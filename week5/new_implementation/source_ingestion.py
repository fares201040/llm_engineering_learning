"""Deterministic discovery and classification of tabular HR source files."""

import csv
from dataclasses import dataclass
from datetime import date, datetime, time
import hashlib
from io import StringIO
import json
from pathlib import Path
from typing import Literal


DOMAIN_RULES = {
    "attendance": {
        "folder": "attendance",
        "required_headers": frozenset({"Employee_ID", "Name", "Date"}),
    }
}

_HEADER_ALIASES = {"OT_value_3": "OT_Value_3"}


@dataclass(frozen=True)
class SourceFile:
    path: Path
    relative_path: str
    sha256: str
    source_format: Literal["xlsx", "csv"]


@dataclass(frozen=True)
class SourceRow:
    row_number: int
    raw_values: tuple[object, ...]
    normalized_record: dict[str, object]


@dataclass(frozen=True)
class SourcePartition:
    source: SourceFile
    name: str
    header_row: int
    raw_headers: tuple[str, ...]
    normalized_headers: tuple[str, ...]
    rows: tuple[SourceRow, ...]


@dataclass(frozen=True)
class PartitionClassification:
    domain: str
    status: Literal["valid", "quarantined", "invalid"]
    reason_code: str | None


@dataclass(frozen=True)
class RawSourceRow:
    raw_row_key: str
    source_path: str
    source_sha256: str
    source_format: str
    partition_name: str
    source_row: int
    domain: str
    status: str
    reason_code: str | None
    payload_hash: str
    payload_json: dict


def discover_sources(
    root: Path,
    xlsx_glob: str,
    csv_glob: str,
) -> tuple[SourceFile, ...]:
    discovered = {}
    for source_format, pattern in (("xlsx", xlsx_glob), ("csv", csv_glob)):
        for path in root.rglob(pattern):
            if not path.is_file() or path.name.startswith("~$"):
                continue
            relative_path = path.relative_to(root).as_posix()
            discovered[relative_path.casefold()] = SourceFile(
                path=path,
                relative_path=relative_path,
                sha256=_file_sha256(path),
                source_format=source_format,
            )
    return tuple(discovered[key] for key in sorted(discovered))


def read_partitions(source: SourceFile) -> tuple[SourcePartition, ...]:
    if source.source_format == "xlsx":
        return _read_xlsx_partitions(source)
    if source.source_format == "csv":
        return (_read_csv_partition(source),)
    raise ValueError(f"Unsupported source format: {source.source_format!r}")


def classify_partition(partition: SourcePartition) -> PartitionClassification:
    normalized = [header for header in partition.normalized_headers if header]
    if "__read_error__" in normalized:
        return PartitionClassification("unknown", "quarantined", "read_error")
    if len(normalized) != len(set(normalized)):
        return PartitionClassification("unknown", "quarantined", "duplicate_headers")

    declared_domain = partition.source.relative_path.split("/", 1)[0].casefold()
    rule = DOMAIN_RULES.get(declared_domain)
    if rule is None:
        return PartitionClassification("unknown", "quarantined", "unrecognized_domain")
    if not rule["required_headers"].issubset(set(normalized)):
        return PartitionClassification("unknown", "quarantined", "schema_mismatch")
    return PartitionClassification(declared_domain, "valid", None)


def build_raw_rows(
    partition: SourcePartition,
    classification: PartitionClassification,
) -> tuple[RawSourceRow, ...]:
    rows = partition.rows
    if not rows and classification.status != "valid":
        rows = (SourceRow(0, (), {}),)

    results = []
    for row in rows:
        payload = {
            "headers": [_json_safe(value) for value in partition.raw_headers],
            "values": [_json_safe(value) for value in row.raw_values],
        }
        payload_hash = _hash_json(payload)
        identity = "\0".join(
            (
                partition.source.relative_path,
                partition.source.sha256,
                partition.name,
                str(row.row_number),
            )
        )
        results.append(
            RawSourceRow(
                raw_row_key=hashlib.sha256(identity.encode("utf-8")).hexdigest(),
                source_path=partition.source.relative_path,
                source_sha256=partition.source.sha256,
                source_format=partition.source.source_format,
                partition_name=partition.name,
                source_row=row.row_number,
                domain=classification.domain,
                status=classification.status,
                reason_code=classification.reason_code,
                payload_hash=payload_hash,
                payload_json=payload,
            )
        )
    return tuple(results)


def normalize_header(value):
    if value is None:
        return ""
    header = str(value).strip()
    if not header:
        return ""
    header = header.replace(" ", "_")
    return _HEADER_ALIASES.get(header, header)


def _file_sha256(path: Path):
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _json_safe(value):
    if isinstance(value, (date, datetime, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _hash_json(value):
    encoded = json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    ).encode("utf-8")
    return hashlib.sha256(encoded).hexdigest()


def _first_nonempty_row(rows):
    for row_number, row in enumerate(rows, start=1):
        values = tuple(row)
        if any(value is not None and str(value).strip() for value in values):
            return row_number, values, rows
    return 0, (), rows


def _partition_from_rows(source, name, rows):
    header_row, raw_headers, remaining_rows = _first_nonempty_row(rows)
    raw_headers = tuple(raw_headers)
    while raw_headers and (raw_headers[-1] is None or not str(raw_headers[-1]).strip()):
        raw_headers = raw_headers[:-1]
    normalized_headers = tuple(normalize_header(value) for value in raw_headers)
    parsed_rows = []
    for row_number, raw_values in enumerate(remaining_rows, start=header_row + 1):
        raw_values = tuple(raw_values)
        if not any(value is not None and str(value).strip() for value in raw_values):
            continue
        normalized_record = {}
        for index, header in enumerate(normalized_headers):
            if not header or index >= len(raw_values):
                continue
            value = _json_safe(raw_values[index])
            if value is not None and value != "":
                normalized_record[header] = value
        parsed_rows.append(SourceRow(row_number, raw_values, normalized_record))
    return SourcePartition(
        source=source,
        name=name,
        header_row=header_row,
        raw_headers=tuple(_json_safe(value) for value in raw_headers),
        normalized_headers=normalized_headers,
        rows=tuple(parsed_rows),
    )


def _read_xlsx_partitions(source):
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(source.path, read_only=True, data_only=True)
    except Exception as exc:
        return (_read_error_partition(source, "XLSX", exc),)
    try:
        return tuple(
            _partition_from_rows(source, sheet.title, sheet.iter_rows(values_only=True))
            for sheet in workbook.worksheets
        )
    finally:
        workbook.close()


def _read_csv_partition(source):
    try:
        text = source.path.read_text(encoding="utf-8-sig")
        sample = text[:8192]
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        return _partition_from_rows(source, "CSV", csv.reader(StringIO(text), dialect))
    except Exception as exc:
        return _read_error_partition(source, "CSV", exc)


def _read_error_partition(source, name, exc):
    message = f"{type(exc).__name__}: {exc}"
    return SourcePartition(
        source=source,
        name=name,
        header_row=0,
        raw_headers=("__read_error__",),
        normalized_headers=("__read_error__",),
        rows=(SourceRow(0, (message,), {"__read_error__": message}),),
    )
