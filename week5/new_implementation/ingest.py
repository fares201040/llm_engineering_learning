from datetime import date, datetime, time as dt_time
from dataclasses import dataclass, replace
from contextlib import nullcontext
from time import perf_counter
from collections import defaultdict
import hashlib
import json
import logging
import math
import os
import tempfile
from typing import Sequence

from openai import OpenAI
from pydantic import BaseModel, ConfigDict, Field, ValidationError
from openpyxl import load_workbook
import tiktoken

try:
    from .attendance_schema import METADATA_FIELDS, SEARCHABLE_FIELDS
    from .chroma_client import create_chroma_client
    from .config import settings
    from .ingestion_state import IngestionLedger
    from .source_ingestion import (
        RawSourceRow,
        SourceFile,
        build_raw_rows,
        classify_partition,
        discover_sources,
        read_partitions,
    )
except ImportError:  # Running ingest.py directly from its directory.
    from attendance_schema import METADATA_FIELDS, SEARCHABLE_FIELDS
    from chroma_client import create_chroma_client
    from config import settings
    from ingestion_state import IngestionLedger
    from source_ingestion import (
        RawSourceRow,
        SourceFile,
        build_raw_rows,
        classify_partition,
        discover_sources,
        read_partitions,
    )


logger = logging.getLogger(__name__)

DB_NAME = str(settings.chroma_db_path)
COLLECTION_NAME = settings.chroma_collection_name

EMBEDDING_MODEL = settings.embedding_model
EMBEDDING_ENCODING = settings.embedding_encoding
EMBEDDING_MAX_TOKENS = settings.embedding_max_tokens
EMBEDDING_BATCH_MAX_TOKENS = settings.embedding_batch_max_tokens
EMBEDDING_BATCH_MAX_ITEMS = settings.embedding_batch_max_items
CHROMA_BATCH_SIZE = settings.chroma_batch_size

KNOWLEDGE_BASE_PATH = settings.knowledge_base_path
JSONL_OUTPUT_PATH = settings.jsonl_output_path
INVALID_JSONL_PATH = settings.invalid_jsonl_path
SOURCE_MANIFEST_PATH = JSONL_OUTPUT_PATH.with_name(
    f"{JSONL_OUTPUT_PATH.stem}.sources.json"
)
SOURCE_XLSX_GLOB = settings.source_xlsx_glob
SOURCE_CSV_GLOB = settings.source_csv_glob

# Improvement 13: optional PostgreSQL + pgvector sink.
# Chroma remains fully supported. PostgreSQL is used only when enabled.
ENABLE_POSTGRES = settings.enable_postgres
POSTGRES_DSN = settings.postgres_dsn
# Plain PostgreSQL does not ship with pgvector. Keep structured SQL ingestion
# useful without it; set ENABLE_PGVECTOR=true only with a pgvector-capable DB.
ENABLE_PGVECTOR = settings.enable_pgvector
PGVECTOR_DIMENSIONS = settings.pgvector_dimensions
POSTGRES_ATTENDANCE_TABLE = settings.postgres_attendance_table
POSTGRES_CHUNKS_TABLE = settings.postgres_chunks_table

# Improvement 15
ENABLE_EMPLOYEE_PERIOD_CHUNKS = settings.enable_employee_period_chunks
ALLOW_EMPTY_SNAPSHOT = settings.allow_empty_snapshot
ALLOW_INVALID_SNAPSHOT = settings.allow_invalid_snapshot
ALLOW_ATTENDANCE_SOURCE_REMOVAL = settings.allow_attendance_source_removal
SOURCE_PARSER_VERSION = 1


class _LazyOpenAI:
    def __init__(self):
        self._client = None

    def __getattr__(self, name):
        if self._client is None:
            self._client = OpenAI()
        return getattr(self._client, name)


openai = _LazyOpenAI()


DATE_FIELDS = {
    "Date",
    "Schedule_From_Date",
    "Schedule_To_Date",
    "Actual_From_Date",
    "Actual_To_Date",
    "From_Date",
    "To_Date",
    # Source export names these *_Time, but sample data contains dates.
    "Pre_OT_Start_Time",
    "Pre_OT_End_Time",
    "Post_OT_Start_Time",
    "Post_OT_End_Time",
}

TIME_FIELDS = {
    "Schedule_From_Time",
    "Schedule_To_Time",
    "Actual_From_Time",
    "Actual_To_Time",
    "From_Time",
    "To_Time",
}

NUMERIC_FIELDS = {
    "Total_Worked_Hrs",
    "Lateness_Hrs",
    "Early_Out_Hrs",
    "Overbreak_Hrs",
    "Regular_Units",
    "Post_OT_hrs",
    "Total_OT",
    "OT_Authorized",
    "OT_Not_Authorized",
    "Leave_Hrs",
    "pre_ot_hrs",
}
NUMERIC_FIELDS.update({f"OT_Value_{i}" for i in range(1, 6)})
NUMERIC_FIELDS.update({f"Allw_Value_{i}" for i in range(1, 11)})

DATE_INPUT_FORMATS = (
    "%Y-%m-%d",
    "%b %d %Y",
    "%b  %d %Y",
    "%m/%d/%Y",
    "%d/%m/%Y",
)

TIME_INPUT_FORMATS = (
    "%H:%M:%S",
    "%H:%M:%S.%f",
    "%H:%M",
    "%I:%M:%S %p",
    "%I:%M %p",
)

DATETIME_INPUT_FORMATS = (
    "%Y-%m-%dT%H:%M:%S",
    "%Y-%m-%d %H:%M:%S",
    "%m/%d/%Y %I:%M:%S %p",
    "%m/%d/%Y %I:%M %p",
    "%m/%d/%Y %H:%M:%S",
)


class AttendanceRecord(BaseModel):
    model_config = ConfigDict(extra="allow")

    Employee_ID: str = Field(min_length=1)
    Name: str = Field(min_length=1)
    Date: str = Field(min_length=10)


class Result(BaseModel):
    page_content: str
    metadata: dict
    record_id: str
    content_hash: str
    chunk_type: str = "attendance_record"


class IngestionStats(BaseModel):
    jsonl_records: int = 0
    invalid_records: int = 0
    attendance_chunks: int = 0
    period_chunks: int = 0
    new_records: int = 0
    changed_records: int = 0
    unchanged_records: int = 0
    deleted_records: int = 0
    duplicate_records: int = 0
    embedding_inputs: int = 0
    embedding_api_batches: int = 0
    chroma_upserts: int = 0
    metadata_updates: int = 0
    postgres_upserts: int = 0
    postgres_attendance_upserts: int = 0
    raw_rows_seen: int = 0
    raw_rows_inserted: int = 0
    unknown_partitions: int = 0
    quarantined_rows: int = 0
    stage_seconds: dict[str, float] = Field(default_factory=dict)
    elapsed_seconds: float = 0.0


@dataclass(frozen=True)
class SourceConversionResult:
    documents: tuple[dict, ...]
    raw_rows: tuple[RawSourceRow, ...]
    partition_summaries: tuple[dict, ...]
    unsafe_reasons: tuple[str, ...]
    invalid_count: int


def _normalize_date_value(value):
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.date().isoformat()

    if isinstance(value, date):
        return value.isoformat()

    text = " ".join(str(value).strip().split())
    if not text:
        return None

    try:
        return datetime.fromisoformat(text).date().isoformat()
    except ValueError:
        pass

    for fmt in DATE_INPUT_FORMATS:
        try:
            return datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            continue

    raise ValueError(f"Invalid date value: {value!r}")


def _normalize_time_value(value):
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.time().replace(microsecond=0).isoformat()

    if isinstance(value, dt_time):
        return value.replace(microsecond=0).isoformat()

    text = str(value).strip()
    if not text:
        return None

    for fmt in TIME_INPUT_FORMATS:
        try:
            return (
                datetime.strptime(text, fmt).time().replace(microsecond=0).isoformat()
            )
        except ValueError:
            continue

    raise ValueError(f"Invalid time value: {value!r}")


def _normalize_numeric_value(value):
    if value is None or value == "":
        return None

    if isinstance(value, bool):
        raise ValueError(f"Boolean is not a valid numeric value: {value!r}")

    if isinstance(value, (int, float)):
        number = float(value)
    else:
        text = str(value).strip().replace(",", "")
        if not text:
            return None
        number = float(text)

    if not math.isfinite(number):
        raise ValueError(f"Non-finite numeric value is not valid: {value!r}")

    if number.is_integer():
        return int(number)
    return number


def _normalize_datetime_value(value):
    if value is None or value == "":
        return None

    if isinstance(value, datetime):
        return value.replace(microsecond=0).isoformat()

    if isinstance(value, date):
        return value.isoformat()

    text = " ".join(str(value).strip().split())
    if not text:
        return None

    try:
        return datetime.fromisoformat(text).replace(microsecond=0).isoformat()
    except ValueError:
        pass

    for fmt in DATETIME_INPUT_FORMATS:
        try:
            return datetime.strptime(text, fmt).replace(microsecond=0).isoformat()
        except ValueError:
            continue

    raise ValueError(f"Invalid datetime value: {value!r}")


def _normalize_json_value(value):
    if value is None:
        return None

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    if isinstance(value, str):
        value = value.strip()
        return value if value else None

    return value


def normalize_attendance_record(record):
    normalized = {}

    for key, value in record.items():
        if key in DATE_FIELDS:
            value = _normalize_date_value(value)
        elif key in TIME_FIELDS:
            value = _normalize_time_value(value)
        elif key in NUMERIC_FIELDS:
            value = _normalize_numeric_value(value)
        elif key == "last_Updated_date":
            value = _normalize_datetime_value(value)
        else:
            value = _normalize_json_value(value)

        if value is not None and value != "":
            normalized[key] = value

    return normalized


def validate_attendance_record(record):
    normalized = normalize_attendance_record(record)
    validated = AttendanceRecord.model_validate(normalized)
    return validated.model_dump(exclude_none=True)


def _normalize_header(value, index):
    if value is None:
        return None

    header = str(value).strip()
    if not header:
        return None

    header = header.replace(" ", "_")

    aliases = {
        "OT_value_3": "OT_Value_3",
    }
    return aliases.get(header, header)


def _trim_trailing_empty_headers(raw_headers):
    headers = list(raw_headers)

    while headers:
        value = headers[-1]
        if value is None or str(value).strip() == "":
            headers.pop()
        else:
            break

    return headers


def _write_invalid_record(handle, record, error, source_file, sheet, excel_row):
    payload = {
        "error": str(error),
        "source_file": source_file,
        "sheet": sheet,
        "excel_row": excel_row,
        "record": record,
    }
    handle.write(json.dumps(payload, ensure_ascii=False, default=str) + "\n")


def _source_excel_files():
    return sorted(
        (
            file
            for file in KNOWLEDGE_BASE_PATH.rglob(SOURCE_XLSX_GLOB)
            if not file.name.startswith("~$")
        ),
        key=lambda file: file.as_posix().casefold(),
    )


def _source_files():
    return discover_sources(
        KNOWLEDGE_BASE_PATH,
        SOURCE_XLSX_GLOB,
        SOURCE_CSV_GLOB,
    )


def _source_signature(source_files):
    """Return a content-based signature for supported source inputs."""
    signature = []

    for source_file in source_files:
        if isinstance(source_file, SourceFile):
            signature.append(
                {
                    "path": source_file.relative_path,
                    "sha256": source_file.sha256,
                    "format": source_file.source_format,
                }
            )
            continue

        digest = hashlib.sha256()
        with open(source_file, "rb") as handle:
            for block in iter(lambda: handle.read(1024 * 1024), b""):
                digest.update(block)

        try:
            display_path = source_file.relative_to(KNOWLEDGE_BASE_PATH).as_posix()
        except ValueError:
            display_path = source_file.resolve().as_posix()

        signature.append(
            {
                "path": display_path,
                "sha256": digest.hexdigest(),
            }
        )

    return signature


def _unsafe_attendance_source_changes(
    previous_partitions,
    current_partitions,
    *,
    allow_removal=False,
):
    """Describe previously valid attendance partitions that are no longer safe."""
    current_by_key = {
        (item.get("source_path"), item.get("name")): item for item in current_partitions
    }
    reasons = []
    for previous in previous_partitions:
        if previous.get("domain") != "attendance" or previous.get("status") != "valid":
            continue
        key = (previous.get("source_path"), previous.get("name"))
        current = current_by_key.get(key)
        if current is None:
            if not allow_removal:
                reasons.append(
                    f"previous attendance partition disappeared: {key[0]}#{key[1]}"
                )
        elif current.get("domain") != "attendance" or current.get("status") != "valid":
            reasons.append(
                f"previous attendance partition is no longer valid: {key[0]}#{key[1]}"
            )
    return reasons


def _jsonl_needs_refresh(source_files=None):
    if not JSONL_OUTPUT_PATH.exists() or not SOURCE_MANIFEST_PATH.exists():
        return True

    try:
        with open(SOURCE_MANIFEST_PATH, "r", encoding="utf-8") as handle:
            manifest = json.load(handle)
    except (OSError, json.JSONDecodeError):
        return True

    if not isinstance(manifest, dict) or not isinstance(
        manifest.get("sources"),
        list,
    ):
        return True

    current_signature = _source_signature(
        source_files if source_files is not None else _source_files()
    )
    if manifest["sources"] != current_signature:
        return True
    if manifest.get("version") != settings.ingestion_format_version:
        return True
    if manifest.get("parser_version") != SOURCE_PARSER_VERSION:
        return True
    outputs = manifest.get("outputs") or {}
    for path, key in (
        (JSONL_OUTPUT_PATH, "jsonl"),
        (INVALID_JSONL_PATH, "invalid_jsonl"),
    ):
        expected = outputs.get(key) or {}
        if not path.exists() or expected.get("sha256") != _file_sha256(path):
            return True
    return False


def _read_source_manifest():
    try:
        with open(SOURCE_MANIFEST_PATH, "r", encoding="utf-8") as handle:
            manifest = json.load(handle)
    except (OSError, json.JSONDecodeError):
        return {}
    return manifest if isinstance(manifest, dict) else {}


def _file_sha256(path):
    digest = hashlib.sha256()
    with open(path, "rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _write_source_manifest(
    source_signature,
    valid_count=None,
    invalid_count=None,
    partition_summaries=None,
):
    SOURCE_MANIFEST_PATH.parent.mkdir(parents=True, exist_ok=True)
    temporary_path = None

    try:
        with tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=SOURCE_MANIFEST_PATH.parent,
            prefix=f".{SOURCE_MANIFEST_PATH.name}.",
            suffix=".tmp",
            delete=False,
        ) as handle:
            temporary_path = handle.name
            json.dump(
                {
                    "version": settings.ingestion_format_version,
                    "parser_version": SOURCE_PARSER_VERSION,
                    "sources": source_signature,
                    "partitions": partition_summaries or [],
                    "outputs": {
                        "jsonl": {
                            "sha256": _file_sha256(JSONL_OUTPUT_PATH),
                            "rows": valid_count,
                        },
                        "invalid_jsonl": {
                            "sha256": _file_sha256(INVALID_JSONL_PATH),
                            "rows": invalid_count,
                        },
                    },
                    "index_schema_version": settings.index_schema_version,
                },
                handle,
                ensure_ascii=False,
                indent=2,
            )
            handle.write("\n")

        os.replace(temporary_path, SOURCE_MANIFEST_PATH)
    finally:
        if temporary_path and os.path.exists(temporary_path):
            os.unlink(temporary_path)


def convert_excel_to_jsonl(ledger=None):
    JSONL_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    INVALID_JSONL_PATH.parent.mkdir(parents=True, exist_ok=True)

    excel_files = _source_excel_files()

    if not excel_files:
        raise FileNotFoundError(f"No .xlsx files found under {KNOWLEDGE_BASE_PATH}")

    source_signature = _source_signature(excel_files)
    ledger = ledger or IngestionLedger(settings.ingestion_state_path)

    written = 0
    invalid = 0

    temporary_paths = []
    refreshed_snapshots = []
    try:
        out = tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=JSONL_OUTPUT_PATH.parent,
            prefix=f".{JSONL_OUTPUT_PATH.name}.",
            suffix=".tmp",
            delete=False,
        )
        invalid_out = tempfile.NamedTemporaryFile(
            mode="w",
            encoding="utf-8",
            dir=INVALID_JSONL_PATH.parent,
            prefix=f".{INVALID_JSONL_PATH.name}.",
            suffix=".tmp",
            delete=False,
        )
        temporary_paths = [out.name, invalid_out.name]
        with out, invalid_out:
            for file, source_entry in zip(excel_files, source_signature):
                source_key = source_entry["path"]
                source_hash = source_entry["sha256"]
                cached = ledger.source_snapshot(source_key, source_hash)
                if cached is not None:
                    for record in cached["valid"]:
                        out.write(
                            json.dumps(record, ensure_ascii=False, default=str) + "\n"
                        )
                    for payload in cached["invalid"]:
                        invalid_out.write(
                            json.dumps(payload, ensure_ascii=False, default=str) + "\n"
                        )
                    written += len(cached["valid"])
                    invalid += len(cached["invalid"])
                    continue

                logger.info("Reading Excel source file=%s", file)
                source_valid = []
                source_invalid = []

                workbook = load_workbook(
                    filename=file,
                    read_only=True,
                    data_only=True,
                )

                try:
                    for sheet in workbook.worksheets:
                        rows = sheet.iter_rows(values_only=True)

                        try:
                            first_row = next(rows)
                        except StopIteration:
                            continue

                        raw_headers = _trim_trailing_empty_headers(first_row)
                        headers = [
                            _normalize_header(value, index)
                            for index, value in enumerate(raw_headers)
                        ]

                        valid_columns = [
                            (index, header)
                            for index, header in enumerate(headers)
                            if header is not None
                        ]

                        if not valid_columns:
                            continue

                        for excel_row_number, row in enumerate(rows, start=2):
                            record = {}

                            for index, key in valid_columns:
                                if index >= len(row):
                                    continue

                                value = _normalize_json_value(row[index])
                                if value is None or value == "":
                                    continue

                                record[key] = value

                            if not record:
                                continue

                            raw_record = dict(record)

                            try:
                                record = validate_attendance_record(record)
                            except (ValidationError, ValueError, TypeError) as exc:
                                invalid += 1
                                invalid_payload = {
                                    "error": str(exc),
                                    "source_file": file.as_posix(),
                                    "sheet": sheet.title,
                                    "excel_row": excel_row_number,
                                    "record": raw_record,
                                }
                                source_invalid.append(invalid_payload)
                                invalid_out.write(
                                    json.dumps(
                                        invalid_payload,
                                        ensure_ascii=False,
                                        default=str,
                                    )
                                    + "\n"
                                )
                                continue

                            record["_source_file"] = file.as_posix()
                            record["_sheet"] = sheet.title
                            record["_excel_row"] = excel_row_number

                            out.write(
                                json.dumps(record, ensure_ascii=False, default=str)
                                + "\n"
                            )
                            source_valid.append(record)
                            written += 1

                finally:
                    workbook.close()
                refreshed_snapshots.append(
                    (source_key, source_hash, source_valid, source_invalid)
                )
            out.flush()
            os.fsync(out.fileno())
            invalid_out.flush()
            os.fsync(invalid_out.fileno())

        os.replace(temporary_paths[0], JSONL_OUTPUT_PATH)
        temporary_paths[0] = None
        os.replace(temporary_paths[1], INVALID_JSONL_PATH)
        temporary_paths[1] = None
        with ledger.writer() as writer:
            writer.replace_source_snapshots(
                refreshed_snapshots,
                [entry["path"] for entry in source_signature],
            )
    finally:
        for temporary_path in temporary_paths:
            if temporary_path and os.path.exists(temporary_path):
                os.unlink(temporary_path)

    _write_source_manifest(source_signature, written, invalid)
    logger.info(
        "Excel conversion complete valid_records=%s invalid_records=%s",
        written,
        invalid,
    )
    return JSONL_OUTPUT_PATH


def convert_sources_to_jsonl(ledger=None):
    previous_partitions = _read_source_manifest().get("partitions") or []
    JSONL_OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    INVALID_JSONL_PATH.parent.mkdir(parents=True, exist_ok=True)

    sources = _source_files()
    if not sources:
        # The source-removal override may approve removing one partition while
        # others remain, but it never authorizes publishing an empty dataset.
        # Preserve the last-safe artifacts and surface the complete removal as
        # an unsafe candidate snapshot.
        removal_reasons = _unsafe_attendance_source_changes(
            previous_partitions,
            [],
            allow_removal=False,
        )
        if removal_reasons:
            return SourceConversionResult(
                documents=(),
                raw_rows=(),
                partition_summaries=(),
                unsafe_reasons=tuple(removal_reasons),
                invalid_count=0,
            )
        raise FileNotFoundError(
            f"No .xlsx or .csv files found under {KNOWLEDGE_BASE_PATH}"
        )

    valid_records = []
    invalid_payloads = []
    raw_rows = []
    summaries = []
    unsafe_reasons = []

    for source in sources:
        logger.info(
            "Reading source file=%s format=%s", source.path, source.source_format
        )
        for partition in read_partitions(source):
            classification = classify_partition(partition)
            partition_raw_rows = list(build_raw_rows(partition, classification))
            summary = {
                "source_path": source.relative_path,
                "name": partition.name,
                "domain": classification.domain,
                "status": classification.status,
                "reason_code": classification.reason_code,
                "row_count": len(partition.rows),
            }
            summaries.append(summary)

            if (
                classification.domain != "attendance"
                or classification.status != "valid"
            ):
                raw_rows.extend(partition_raw_rows)
                continue

            for source_row, raw_row in zip(partition.rows, partition_raw_rows):
                raw_record = dict(source_row.normalized_record)
                try:
                    record = validate_attendance_record(raw_record)
                except (ValidationError, ValueError, TypeError) as exc:
                    invalid_payloads.append(
                        {
                            "error": str(exc),
                            "source_file": source.path.as_posix(),
                            "sheet": partition.name,
                            "excel_row": source_row.row_number,
                            "record": raw_record,
                        }
                    )
                    raw_rows.append(
                        replace(
                            raw_row,
                            status="invalid",
                            reason_code="validation_error",
                        )
                    )
                    continue

                record["_source_file"] = source.path.as_posix()
                record["_sheet"] = partition.name
                record["_excel_row"] = source_row.row_number
                record["_source_format"] = source.source_format
                record["_raw_row_key"] = raw_row.raw_row_key
                valid_records.append(record)
                raw_rows.append(raw_row)

    unsafe_reasons.extend(
        _unsafe_attendance_source_changes(
            previous_partitions,
            summaries,
            allow_removal=ALLOW_ATTENDANCE_SOURCE_REMOVAL,
        )
    )

    if not unsafe_reasons:
        temporary_paths = []
        try:
            output = tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                dir=JSONL_OUTPUT_PATH.parent,
                prefix=f".{JSONL_OUTPUT_PATH.name}.",
                suffix=".tmp",
                delete=False,
            )
            invalid_output = tempfile.NamedTemporaryFile(
                mode="w",
                encoding="utf-8",
                dir=INVALID_JSONL_PATH.parent,
                prefix=f".{INVALID_JSONL_PATH.name}.",
                suffix=".tmp",
                delete=False,
            )
            temporary_paths = [output.name, invalid_output.name]
            with output, invalid_output:
                for record in valid_records:
                    output.write(
                        json.dumps(record, ensure_ascii=False, default=str) + "\n"
                    )
                for payload in invalid_payloads:
                    invalid_output.write(
                        json.dumps(payload, ensure_ascii=False, default=str) + "\n"
                    )
                output.flush()
                os.fsync(output.fileno())
                invalid_output.flush()
                os.fsync(invalid_output.fileno())

            os.replace(temporary_paths[0], JSONL_OUTPUT_PATH)
            temporary_paths[0] = None
            os.replace(temporary_paths[1], INVALID_JSONL_PATH)
            temporary_paths[1] = None
        finally:
            for temporary_path in temporary_paths:
                if temporary_path and os.path.exists(temporary_path):
                    os.unlink(temporary_path)

    signature = [
        {
            "path": source.relative_path,
            "sha256": source.sha256,
            "format": source.source_format,
        }
        for source in sources
    ]
    if not unsafe_reasons:
        _write_source_manifest(
            signature,
            len(valid_records),
            len(invalid_payloads),
            summaries,
        )
    if ledger is not None and not unsafe_reasons:
        refreshed_snapshots = []
        for source in sources:
            source_valid = [
                record
                for record in valid_records
                if record.get("_source_file") == source.path.as_posix()
            ]
            source_invalid = [
                item
                for item in invalid_payloads
                if item.get("source_file") == source.path.as_posix()
            ]
            source_summaries = [
                item
                for item in summaries
                if item["source_path"] == source.relative_path
            ]
            source_raw_count = sum(
                1 for row in raw_rows if row.source_path == source.relative_path
            )
            refreshed_snapshots.append(
                (
                    source.relative_path,
                    source.sha256,
                    source_valid,
                    source_invalid,
                    SOURCE_PARSER_VERSION,
                    source_summaries,
                    source_raw_count,
                )
            )
        with ledger.writer() as writer:
            writer.replace_source_snapshots(
                refreshed_snapshots,
                [source.relative_path for source in sources],
            )

    documents = tuple(
        {
            "type": "attendance",
            "source": record["_source_file"],
            "sheet": record.get("_sheet"),
            "excel_row": record.get("_excel_row"),
            "jsonl_line": line_number,
            "record": dict(record),
        }
        for line_number, record in enumerate(valid_records, start=1)
    )
    return SourceConversionResult(
        documents=documents,
        raw_rows=tuple(raw_rows),
        partition_summaries=tuple(summaries),
        unsafe_reasons=tuple(unsafe_reasons),
        invalid_count=len(invalid_payloads),
    )


def _documents_from_jsonl():
    """Read and validate the current attendance projection without mutating it."""
    documents = []

    with open(JSONL_OUTPUT_PATH, "r", encoding="utf-8") as f:
        for line_number, line in enumerate(f, start=1):
            line = line.strip()
            if not line:
                continue

            try:
                record = json.loads(line)
            except json.JSONDecodeError as exc:
                raise ValueError(
                    f"Invalid JSONL at line {line_number} in {JSONL_OUTPUT_PATH}: {exc}"
                ) from exc

            if not isinstance(record, dict):
                raise ValueError(
                    f"JSONL line {line_number} must contain a JSON object."
                )

            trace = {
                key: record.get(key)
                for key in (
                    "_source_file",
                    "_sheet",
                    "_excel_row",
                    "_source_format",
                    "_raw_row_key",
                )
            }
            business_record = {
                key: value for key, value in record.items() if not key.startswith("_")
            }

            try:
                business_record = validate_attendance_record(business_record)
            except (ValidationError, ValueError, TypeError) as exc:
                logger.warning(
                    "Skipping invalid JSONL record line=%s error=%s",
                    line_number,
                    exc,
                )
                continue

            record = dict(business_record)
            for key, value in trace.items():
                if value is not None:
                    record[key] = value

            documents.append(
                {
                    "type": "attendance",
                    "source": record.get(
                        "_source_file",
                        JSONL_OUTPUT_PATH.as_posix(),
                    ),
                    "sheet": record.get("_sheet"),
                    "excel_row": record.get("_excel_row"),
                    "jsonl_line": line_number,
                    "record": record,
                }
            )
    return documents


def _raw_postgres_coverage_complete(sources, partitions):
    if not ENABLE_POSTGRES:
        return True
    if not POSTGRES_DSN:
        raise ValueError("ENABLE_POSTGRES is true but POSTGRES_DSN is empty.")
    try:
        import psycopg
    except ImportError as exc:
        raise RuntimeError(
            'PostgreSQL is enabled. Install psycopg with: uv add "psycopg[binary]"'
        ) from exc

    expected_by_source = defaultdict(int)
    for partition in partitions:
        row_count = int(partition.get("row_count") or 0)
        if partition.get("status") != "valid" and row_count == 0:
            row_count = 1  # Row-zero quarantine sentinel.
        expected_by_source[partition.get("source_path")] += row_count

    with psycopg.connect(POSTGRES_DSN) as connection:
        with connection.cursor() as cursor:
            cursor.execute("SELECT to_regclass('private_ingestion.raw_source_rows')")
            if cursor.fetchone()[0] is None:
                return False
            for source in sources:
                cursor.execute(
                    """
                    SELECT COUNT(*)
                    FROM private_ingestion.raw_source_rows
                    WHERE source_path = %s AND source_sha256 = %s
                    """,
                    (source.relative_path, source.sha256),
                )
                actual = cursor.fetchone()[0]
                if actual != expected_by_source[source.relative_path]:
                    return False
    return True


def load_source_snapshot(ledger=None):
    """Return the current typed projection plus raw rows from any required parse."""
    source_files = _source_files()
    manifest = _read_source_manifest()
    if _jsonl_needs_refresh(source_files) or not _raw_postgres_coverage_complete(
        source_files,
        manifest.get("partitions") or (),
    ):
        return convert_sources_to_jsonl(ledger=ledger)

    return SourceConversionResult(
        documents=tuple(_documents_from_jsonl()),
        raw_rows=(),
        partition_summaries=tuple(manifest.get("partitions") or ()),
        unsafe_reasons=(),
        invalid_count=_count_invalid_records(),
    )


def fetch_jsonl_documents(ledger=None):
    return list(load_source_snapshot(ledger=ledger).documents)


def fetch_documents(ledger=None):
    return fetch_jsonl_documents(ledger=ledger)


LEGACY_SEARCH_TEXT_FIELDS = [
    "Employee_ID",
    "Name",
    "Organization_Unit",
    "Work_Location",
    "Department",
    "Position",
    "Job",
    "Grade",
    "Date",
    "Day",
    "Day_Type",
    "Holiday_Type",
    "Shift",
    "Status",
    "Exception",
    "Schedule_From_Time",
    "Schedule_To_Time",
    "Actual_From_Time",
    "Actual_To_Time",
    "Total_Worked_Hrs",
    "Lateness_Hrs",
    "Early_Out_Hrs",
    "Overbreak_Hrs",
    "Regular_Units",
    "Late_In_Reason",
    "Early_Out_Reason",
    "Employee_Remarks",
    "Approvers_remarks",
    "pre_ot_hrs",
    "Post_OT_hrs",
    "Total_OT",
    "OT_Authorized",
    "OT_Not_Authorized",
    "OT_Type_1",
    "OT_Value_1",
    "OT_Type_2",
    "OT_Value_2",
    "OT_Type_3",
    "OT_Value_3",
    "OT_Type_4",
    "OT_Value_4",
    "OT_Type_5",
    "OT_Value_5",
    "Leave_Type",
    "Leave_Hrs",
]

LEGACY_METADATA_FIELDS = [
    "Employee_ID",
    "Name",
    "Organization_Unit",
    "Country",
    "Work_Location",
    "Department",
    "Position",
    "Job",
    "Gradeset",
    "Grade",
    "Date",
    "Day",
    "Day_Type",
    "Holiday_Type",
    "Shift",
    "Status",
    "Exception",
    "Total_Worked_Hrs",
    "Lateness_Hrs",
    "Early_Out_Hrs",
    "Overbreak_Hrs",
    "Regular_Units",
    "pre_ot_hrs",
    "Post_OT_hrs",
    "Total_OT",
    "OT_Authorized",
    "OT_Not_Authorized",
    "OT_Type_1",
    "OT_Value_1",
    "OT_Type_2",
    "OT_Value_2",
    "OT_Type_3",
    "OT_Value_3",
    "OT_Type_4",
    "OT_Value_4",
    "OT_Type_5",
    "OT_Value_5",
    "Leave_Type",
    "Leave_Hrs",
]

# Compatibility aliases now derived from the canonical executable registry.
SEARCH_TEXT_FIELDS = list(SEARCHABLE_FIELDS)
METADATA_FIELDS = list(METADATA_FIELDS)


def _record_to_search_text(record):
    lines = []

    for key in SEARCH_TEXT_FIELDS:
        value = record.get(key)
        if value is None or value == "":
            continue
        lines.append(f"{key}: {value}")

    return "\n".join(lines).strip()


def _metadata_scalar(value):
    if isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def _record_to_metadata(document):
    record = document["record"]

    metadata = {
        "domain": "attendance",
        "type": document["type"],
        "source": document["source"],
        "sheet": document.get("sheet"),
        "excel_row": document.get("excel_row"),
        "jsonl_line": document.get("jsonl_line"),
        "chunk_type": "attendance_record",
    }
    if record.get("_raw_row_key"):
        metadata["raw_row_key"] = str(record["_raw_row_key"])

    for field in METADATA_FIELDS:
        value = record.get(field)
        if value is None or value == "":
            continue
        metadata[field] = _metadata_scalar(value)

    return {key: value for key, value in metadata.items() if value is not None}


def _canonical_json(value):
    return json.dumps(
        value,
        ensure_ascii=False,
        sort_keys=True,
        separators=(",", ":"),
        default=str,
    )


def _business_record(record):
    """Return business fields only; exclude source/trace fields starting with '_'."""
    return {key: value for key, value in record.items() if not key.startswith("_")}


def _content_hash(value):
    """
    Improvement 11/12:
    Stable hash used to detect whether a record changed.
    """
    return hashlib.sha256(_canonical_json(value).encode("utf-8")).hexdigest()


def _slug(value):
    text = str(value or "").strip().lower()
    cleaned = "".join(ch if ch.isalnum() else "-" for ch in text)
    return "-".join(part for part in cleaned.split("-") if part) or "na"


def _make_record_id(record):
    """
    Improvement 8:
    Stable logical attendance-record ID.

    The key intentionally uses source-independent business values so the same
    attendance record receives the same ID after re-ingestion.
    """
    identity = {
        "Employee_ID": record.get("Employee_ID"),
        "Date": record.get("Date"),
        "Shift": record.get("Shift"),
        "Schedule_From_Time": record.get("Schedule_From_Time"),
        "Schedule_To_Time": record.get("Schedule_To_Time"),
    }

    digest = hashlib.sha1(_canonical_json(identity).encode("utf-8")).hexdigest()[:12]

    return (
        f"attendance:"
        f"{_slug(record.get('Employee_ID'))}:"
        f"{_slug(record.get('Date'))}:"
        f"{digest}"
    )


def _record_update_sort_key(record):
    """
    Prefer the newest duplicate record when last_Updated_date is available.
    ISO timestamps sort correctly as strings. Missing timestamps sort first.
    """
    return str(record.get("last_Updated_date") or "")


def canonicalize_documents(documents, stats=None):
    """Select one authoritative source document per logical attendance record."""
    if stats is None:
        stats = IngestionStats()

    selected = {}
    hashes = {}
    for document in documents:
        record = document["record"]
        record_id = _make_record_id(record)
        content_hash = _content_hash(_business_record(record))
        previous = selected.get(record_id)
        if previous is None:
            selected[record_id] = document
            hashes[record_id] = content_hash
            continue

        stats.duplicate_records += 1
        if hashes[record_id] == content_hash:
            continue
        if _record_update_sort_key(record) >= _record_update_sort_key(
            previous["record"]
        ):
            selected[record_id] = document
            hashes[record_id] = content_hash

    return list(selected.values())


def create_record_chunks(documents, stats=None):
    """
    Improvements 8, 11 and 12:
    - one stable logical ID per attendance record
    - exact duplicate records are stored only once
    - conflicting duplicates keep the newest source version
    """
    if stats is None:
        stats = IngestionStats()

    chunks = []
    for document in canonicalize_documents(documents, stats):
        record = document["record"]
        page_content = _record_to_search_text(record)

        if not page_content:
            continue

        record_id = _make_record_id(record)
        content_hash = _content_hash(_business_record(record))

        metadata = _record_to_metadata(document)
        metadata["record_id"] = record_id
        metadata["content_hash"] = content_hash
        chunks.append(
            Result(
                page_content=page_content,
                metadata=metadata,
                record_id=record_id,
                content_hash=content_hash,
                chunk_type="attendance_record",
            )
        )

    return chunks


def _month_period(date_text):
    dt = datetime.strptime(date_text, "%Y-%m-%d")
    return dt.strftime("%Y-%m")


def create_employee_period_chunks(documents):
    """
    Improvement 15:
    Create one deterministic monthly summary chunk per employee.

    This is generated in Python, not by an LLM.
    """
    grouped = defaultdict(list)

    for document in documents:
        record = document["record"]
        employee_id = record.get("Employee_ID")
        record_date = record.get("Date")

        if not employee_id or not record_date:
            continue

        grouped[(employee_id, _month_period(record_date))].append(record)

    chunks = []

    for (employee_id, period), records in grouped.items():
        records.sort(key=lambda item: item.get("Date", ""))

        first = records[0]
        name = first.get("Name", "")
        department = first.get("Department", "")
        position = first.get("Position", "")

        lines = [
            f"Employee_ID: {employee_id}",
            f"Name: {name}",
            f"Period: {period}",
        ]

        if department:
            lines.append(f"Department: {department}")
        if position:
            lines.append(f"Position: {position}")

        lines.append("Attendance records:")

        for record in records:
            details = [record.get("Date", "")]

            for field, label in [
                ("Shift", "Shift"),
                ("Exception", "Exception"),
                ("Total_Worked_Hrs", "Worked"),
                ("Lateness_Hrs", "Late"),
                ("Early_Out_Hrs", "EarlyOut"),
                ("Total_OT", "OT"),
                ("Leave_Type", "Leave"),
            ]:
                value = record.get(field)
                if value is not None and value != "":
                    details.append(f"{label}={value}")

            lines.append(" | ".join(details))

        page_content = "\n".join(lines)

        summary_identity = {
            "Employee_ID": employee_id,
            "Period": period,
            "chunk_type": "employee_period",
        }
        record_id = (
            f"period:{_slug(employee_id)}:{period}:"
            f"{hashlib.sha1(_canonical_json(summary_identity).encode()).hexdigest()[:10]}"
        )
        content_hash = _content_hash(page_content)

        metadata = {
            "domain": "attendance",
            "type": "attendance",
            "chunk_type": "employee_period",
            "record_id": record_id,
            "content_hash": content_hash,
            "Employee_ID": employee_id,
            "Name": name,
            "Period": period,
            "Department": department or "Unknown",
            "records_count": len(records),
        }

        chunks.append(
            Result(
                page_content=page_content,
                metadata=metadata,
                record_id=record_id,
                content_hash=content_hash,
                chunk_type="employee_period",
            )
        )

    return chunks


SPLIT_IDENTITY_FIELDS = [
    "Employee_ID",
    "Name",
    "Date",
    "Period",
    "Department",
    "Position",
    "Shift",
]


def _build_split_identity_prefix(metadata):
    lines = []

    for field in SPLIT_IDENTITY_FIELDS:
        value = metadata.get(field)
        if value is None or value == "":
            continue
        lines.append(f"{field}: {value}")

    return "\n".join(lines).strip()


def _split_text_for_embedding(
    text,
    identity_prefix,
    encoding,
    max_tokens=None,
):
    if max_tokens is None:
        max_tokens = EMBEDDING_MAX_TOKENS

    text_tokens = encoding.encode(text)

    if len(text_tokens) <= max_tokens:
        return [(text, len(text_tokens))]

    prefix_text = identity_prefix + "\n\n" if identity_prefix else ""
    prefix_tokens = encoding.encode(prefix_text)

    body_max_tokens = max_tokens - len(prefix_tokens)

    if body_max_tokens <= 0:
        raise ValueError(
            "The split identity prefix is too large to fit inside "
            f"the {max_tokens}-token embedding limit."
        )

    parts = []

    for start in range(0, len(text_tokens), body_max_tokens):
        body_tokens = text_tokens[start : start + body_max_tokens]
        part_text = prefix_text + encoding.decode(body_tokens)
        part_token_count = len(encoding.encode(part_text))

        if part_token_count > max_tokens:
            raise ValueError(
                f"Generated embedding part has {part_token_count} tokens, "
                f"above the {max_tokens}-token limit."
            )

        parts.append((part_text, part_token_count))

    return parts


def _prepare_embedding_items(chunks):
    encoding = tiktoken.get_encoding(EMBEDDING_ENCODING)
    items = []

    for chunk_index, chunk in enumerate(chunks):
        identity_prefix = _build_split_identity_prefix(chunk.metadata)

        parts = _split_text_for_embedding(
            text=chunk.page_content,
            identity_prefix=identity_prefix,
            encoding=encoding,
        )

        for part_index, (part_text, token_count) in enumerate(parts, start=1):
            part_id = (
                chunk.record_id
                if len(parts) == 1
                else f"{chunk.record_id}:part-{part_index}"
            )

            metadata = dict(chunk.metadata)
            metadata["embedding_part"] = part_index
            metadata["embedding_parts_total"] = len(parts)
            metadata["part_id"] = part_id
            embedding_input_hash = _content_hash(
                {
                    "model": EMBEDDING_MODEL,
                    "index_schema_version": settings.index_schema_version,
                    "text": part_text,
                }
            )
            metadata_hash = _content_hash(metadata)
            metadata["embedding_input_hash"] = embedding_input_hash
            metadata["metadata_hash"] = metadata_hash

            items.append(
                {
                    "id": part_id,
                    "record_id": chunk.record_id,
                    "content_hash": chunk.content_hash,
                    "chunk_type": chunk.chunk_type,
                    "text": part_text,
                    "metadata": metadata,
                    "token_count": token_count,
                    "embedding_input_hash": embedding_input_hash,
                    "metadata_hash": metadata_hash,
                }
            )

    return items


def _iter_embedding_batches(items):
    batch = []
    batch_tokens = 0

    for item in items:
        exceeds_tokens = (
            bool(batch)
            and batch_tokens + item["token_count"] > EMBEDDING_BATCH_MAX_TOKENS
        )
        exceeds_items = len(batch) >= EMBEDDING_BATCH_MAX_ITEMS

        if exceeds_tokens or exceeds_items:
            yield batch, batch_tokens
            batch = []
            batch_tokens = 0

        batch.append(item)
        batch_tokens += item["token_count"]

    if batch:
        yield batch, batch_tokens


def _existing_chroma_state(collection):
    """
    Improvement 11/12:
    Return existing Chroma rows grouped by logical record_id.
    """
    existing = collection.get(include=["metadatas"])

    state = defaultdict(list)

    for item_id, metadata in zip(
        existing.get("ids", []),
        existing.get("metadatas", []),
    ):
        metadata = metadata or {}
        record_id = metadata.get("record_id")

        if not record_id:
            # Older DB rows from before stable IDs are intentionally considered stale.
            record_id = f"legacy:{item_id}"

        state[record_id].append(
            {
                "id": item_id,
                "content_hash": metadata.get("content_hash"),
                "embedding_input_hash": metadata.get("embedding_input_hash"),
                "metadata_hash": metadata.get("metadata_hash"),
            }
        )

    return state


def _embed_items(items, stats):
    vectors_by_id = {}

    batches = list(_iter_embedding_batches(items))
    stats.embedding_api_batches += len(batches)

    for batch_number, (batch, batch_tokens) in enumerate(batches, start=1):
        texts = [item["text"] for item in batch]

        logger.info(
            "Embedding batch=%s/%s input_count=%s token_count=%s",
            batch_number,
            len(batches),
            len(texts),
            batch_tokens,
        )

        response = openai.embeddings.create(
            model=EMBEDDING_MODEL,
            input=texts,
        )

        if len(response.data) != len(batch):
            raise RuntimeError(
                f"Embedding API returned {len(response.data)} vectors "
                f"for {len(batch)} inputs."
            )

        for item, embedding in zip(batch, response.data):
            vectors_by_id[item["id"]] = embedding.embedding

    return vectors_by_id


def sync_embeddings_to_chroma(chunks, stats=None, ledger=None, generation=None):
    """
    Improvements 8-12:
    - stable IDs
    - duplicate detection
    - incremental embedding/upsert
    - unchanged records skipped
    - deleted source records removed
    - batched embeddings retained
    """
    if stats is None:
        stats = IngestionStats()

    items = _prepare_embedding_items(chunks)
    chroma = create_chroma_client()
    collection = chroma.get_or_create_collection(COLLECTION_NAME)
    existing_state = _existing_chroma_state(collection)

    current_by_record = defaultdict(list)
    for item in items:
        current_by_record[item["record_id"]].append(item)

    desired_record_ids = set(current_by_record)
    existing_record_ids = set(existing_state)

    stale_record_ids = existing_record_ids - desired_record_ids
    stale_item_ids = [
        row["id"] for record_id in stale_record_ids for row in existing_state[record_id]
    ]
    stats.deleted_records = len(stale_record_ids)

    changed_or_new_items = []
    metadata_only_items = []
    obsolete_changed_item_ids = []

    for record_id, record_items in current_by_record.items():
        existing_rows = existing_state.get(record_id)

        if not existing_rows:
            stats.new_records += 1
            changed_or_new_items.extend(record_items)
            continue

        current_ids = {item["id"] for item in record_items}
        existing_ids = {row["id"] for row in existing_rows}

        if existing_ids != current_ids:
            stats.changed_records += 1
            obsolete_changed_item_ids.extend(existing_ids - current_ids)
            changed_or_new_items.extend(record_items)
            continue

        existing_by_id = {row["id"]: row for row in existing_rows}
        if any(
            existing_by_id[item["id"]].get("embedding_input_hash")
            != item["embedding_input_hash"]
            for item in record_items
        ):
            stats.changed_records += 1
            changed_or_new_items.extend(record_items)
            continue

        changed_metadata = [
            item
            for item in record_items
            if existing_by_id[item["id"]].get("metadata_hash") != item["metadata_hash"]
        ]
        if changed_metadata:
            stats.changed_records += 1
            metadata_only_items.extend(changed_metadata)
        else:
            stats.unchanged_records += 1

    for start in range(0, len(metadata_only_items), CHROMA_BATCH_SIZE):
        batch = metadata_only_items[start : start + CHROMA_BATCH_SIZE]
        collection.update(
            ids=[item["id"] for item in batch],
            metadatas=[item["metadata"] for item in batch],
        )
        stats.metadata_updates += len(batch)

    stats.embedding_inputs = len(changed_or_new_items)

    if not changed_or_new_items:
        for start in range(0, len(stale_item_ids), CHROMA_BATCH_SIZE):
            collection.delete(ids=stale_item_ids[start : start + CHROMA_BATCH_SIZE])
        logger.info("Chroma sync has no new_or_changed_records=true")
        return stats

    # Stream embedding -> durable Chroma upsert -> checkpoint. This bounds
    # memory and permits a failed run to resume from its first incomplete batch.
    for batch_number, (embedding_batch, _tokens) in enumerate(
        _iter_embedding_batches(changed_or_new_items), start=1
    ):
        checkpoint_key = f"batch-{batch_number}"
        if (
            ledger is not None
            and generation is not None
            and ledger.is_checkpoint_complete(
                "chroma",
                checkpoint_key,
                generation,
                payload_hash=(
                    payload_hash := _content_hash(
                        [
                            {
                                "id": item["id"],
                                "embedding_input_hash": item["embedding_input_hash"],
                                "metadata_hash": item["metadata_hash"],
                            }
                            for item in embedding_batch
                        ]
                    )
                ),
            )
        ):
            stored = collection.get(
                ids=[item["id"] for item in embedding_batch], include=["metadatas"]
            )
            stored_hashes = {
                item_id: (
                    metadata.get("embedding_input_hash"),
                    metadata.get("metadata_hash"),
                )
                for item_id, metadata in zip(
                    stored.get("ids") or [], stored.get("metadatas") or []
                )
            }
            expected_hashes = {
                item["id"]: (item["embedding_input_hash"], item["metadata_hash"])
                for item in embedding_batch
            }
            if stored_hashes == expected_hashes:
                continue
        else:
            payload_hash = _content_hash(
                [
                    {
                        "id": item["id"],
                        "embedding_input_hash": item["embedding_input_hash"],
                        "metadata_hash": item["metadata_hash"],
                    }
                    for item in embedding_batch
                ]
            )

        vectors_by_id = _embed_items(embedding_batch, stats)
        for start in range(0, len(embedding_batch), CHROMA_BATCH_SIZE):
            batch = embedding_batch[start : start + CHROMA_BATCH_SIZE]
            collection.upsert(
                ids=[item["id"] for item in batch],
                embeddings=[vectors_by_id[item["id"]] for item in batch],
                documents=[item["text"] for item in batch],
                metadatas=[item["metadata"] for item in batch],
            )
            stats.chroma_upserts += len(batch)
        if ledger is not None and generation is not None:
            with ledger.writer() as writer:
                writer.checkpoint(
                    "chroma", checkpoint_key, generation, payload_hash=payload_hash
                )

    # Remove stale records and superseded split parts only after replacement
    # embeddings have been created and every upsert has succeeded.
    ids_to_delete = stale_item_ids + obsolete_changed_item_ids
    for start in range(0, len(ids_to_delete), CHROMA_BATCH_SIZE):
        collection.delete(ids=ids_to_delete[start : start + CHROMA_BATCH_SIZE])

    logger.info(
        "Chroma sync complete count=%s new=%s changed=%s unchanged=%s deleted=%s",
        collection.count(),
        stats.new_records,
        stats.changed_records,
        stats.unchanged_records,
        stats.deleted_records,
    )

    return stats


def _postgres_attendance_values(document):
    """
    Improvement 16:
    Convert one validated JSONL document into typed PostgreSQL values.

    The complete normalized record is also kept in JSONB so PostgreSQL
    becomes the structured source of truth rather than only a vector mirror.
    """
    record = document["record"]
    record_id = _make_record_id(record)
    business_record = _business_record(record)

    return {
        "record_id": record_id,
        "content_hash": _content_hash(business_record),
        "employee_id": record.get("Employee_ID"),
        "name": record.get("Name"),
        "organization_unit": record.get("Organization_Unit"),
        "country": record.get("Country"),
        "work_location": record.get("Work_Location"),
        "department": record.get("Department"),
        "position": record.get("Position"),
        "job": record.get("Job"),
        "gradeset": record.get("Gradeset"),
        "grade": record.get("Grade"),
        "attendance_date": record.get("Date"),
        "day": record.get("Day"),
        "day_type": record.get("Day_Type"),
        "holiday_type": record.get("Holiday_Type"),
        "shift": record.get("Shift"),
        "status": record.get("Status"),
        "exception": record.get("Exception"),
        "total_worked_hrs": record.get("Total_Worked_Hrs"),
        "lateness_hrs": record.get("Lateness_Hrs"),
        "early_out_hrs": record.get("Early_Out_Hrs"),
        "overbreak_hrs": record.get("Overbreak_Hrs"),
        "regular_units": record.get("Regular_Units"),
        "pre_ot_hrs": record.get("pre_ot_hrs"),
        "post_ot_hrs": record.get("Post_OT_hrs"),
        "total_ot": record.get("Total_OT"),
        "ot_authorized": record.get("OT_Authorized"),
        "ot_not_authorized": record.get("OT_Not_Authorized"),
        "leave_type": record.get("Leave_Type"),
        "leave_hrs": record.get("Leave_Hrs"),
        "last_updated_at": record.get("last_Updated_date"),
        "source_file": document.get("source"),
        "source_sheet": document.get("sheet"),
        "source_excel_row": document.get("excel_row"),
        "source_jsonl_line": document.get("jsonl_line"),
        "search_text": _record_to_search_text(record),
        "record_json": business_record,
        "raw_row_key": record.get("_raw_row_key"),
    }


def store_raw_rows_to_postgres(
    raw_rows: Sequence[RawSourceRow],
    *,
    connection=None,
) -> int:
    if not raw_rows:
        return 0
    if connection is None and not ENABLE_POSTGRES:
        return 0
    if connection is None and not POSTGRES_DSN:
        raise ValueError("ENABLE_POSTGRES is true but POSTGRES_DSN is empty.")

    def execute(conn):
        inserted = 0
        with conn.cursor() as cur:
            cur.execute("CREATE SCHEMA IF NOT EXISTS private_ingestion")
            cur.execute(
                """
                CREATE TABLE IF NOT EXISTS private_ingestion.raw_source_rows (
                    raw_row_key TEXT PRIMARY KEY,
                    source_path TEXT NOT NULL,
                    source_sha256 TEXT NOT NULL,
                    source_format TEXT NOT NULL
                        CHECK (source_format IN ('xlsx', 'csv')),
                    partition_name TEXT NOT NULL,
                    source_row INTEGER NOT NULL,
                    domain TEXT NOT NULL,
                    status TEXT NOT NULL
                        CHECK (status IN ('valid', 'quarantined', 'invalid')),
                    reason_code TEXT,
                    payload_hash TEXT NOT NULL,
                    payload_json JSONB NOT NULL,
                    ingested_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                    UNIQUE (source_path, source_sha256, partition_name, source_row)
                )
                """
            )
            cur.execute(
                "CREATE INDEX IF NOT EXISTS idx_raw_source_revision "
                "ON private_ingestion.raw_source_rows "
                "(source_path, source_sha256)"
            )
            cur.execute(
                "CREATE INDEX IF NOT EXISTS idx_raw_source_status "
                "ON private_ingestion.raw_source_rows (domain, status)"
            )
            cur.execute("REVOKE ALL ON SCHEMA private_ingestion FROM PUBLIC")
            cur.execute("REVOKE ALL ON private_ingestion.raw_source_rows FROM PUBLIC")
            for row in raw_rows:
                cur.execute(
                    """
                    INSERT INTO private_ingestion.raw_source_rows (
                        raw_row_key, source_path, source_sha256, source_format,
                        partition_name, source_row, domain, status, reason_code,
                        payload_hash, payload_json
                    ) VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
                    ON CONFLICT (raw_row_key) DO NOTHING
                    """,
                    (
                        row.raw_row_key,
                        row.source_path,
                        row.source_sha256,
                        row.source_format,
                        row.partition_name,
                        row.source_row,
                        row.domain,
                        row.status,
                        row.reason_code,
                        row.payload_hash,
                        json.dumps(row.payload_json, ensure_ascii=False, default=str),
                    ),
                )
                inserted += max(cur.rowcount, 0)
        return inserted

    if connection is not None:
        return execute(connection)

    try:
        import psycopg
    except ImportError as exc:
        raise RuntimeError(
            'PostgreSQL is enabled. Install psycopg with: uv add "psycopg[binary]"'
        ) from exc
    with psycopg.connect(POSTGRES_DSN) as conn:
        return execute(conn)


def publish_postgres_snapshot(
    documents: Sequence[dict],
    raw_rows: Sequence[RawSourceRow],
    generation: str,
    stats: IngestionStats | None = None,
    *,
    connection=None,
) -> IngestionStats:
    del generation
    stats = stats or IngestionStats()
    stats.raw_rows_seen += len(raw_rows)

    if connection is None and not ENABLE_POSTGRES:
        return stats
    if connection is None and not POSTGRES_DSN:
        raise ValueError("ENABLE_POSTGRES is true but POSTGRES_DSN is empty.")

    def publish(conn):
        stats.raw_rows_inserted += store_raw_rows_to_postgres(
            raw_rows,
            connection=conn,
        )
        sync_attendance_records_to_postgres(
            documents,
            stats,
            connection=conn,
        )

    if connection is not None:
        publish(connection)
        return stats

    try:
        import psycopg
    except ImportError as exc:
        raise RuntimeError(
            'PostgreSQL is enabled. Install psycopg with: uv add "psycopg[binary]"'
        ) from exc
    with psycopg.connect(POSTGRES_DSN) as conn:
        publish(conn)
    return stats


def sync_attendance_records_to_postgres(documents, stats=None, *, connection=None):
    """
    Improvement 16:
    Synchronize full structured attendance records into PostgreSQL.

    PostgreSQL becomes the authoritative structured query store when
    ENABLE_POSTGRES=true. Chroma remains available for fallback/vector use.
    """
    if stats is None:
        stats = IngestionStats()

    if not ENABLE_POSTGRES:
        return stats

    if not POSTGRES_DSN:
        raise ValueError("ENABLE_POSTGRES is true but POSTGRES_DSN is empty.")

    psycopg = None
    if connection is None:
        try:
            import psycopg
        except ImportError as exc:
            raise RuntimeError(
                'PostgreSQL is enabled. Install psycopg with: uv add "psycopg[binary]"'
            ) from exc

    rows_by_id = {}

    # Reuse the same newest-version rule as Chroma duplicate handling.
    for document in documents:
        values = _postgres_attendance_values(document)
        previous = rows_by_id.get(values["record_id"])

        if previous is None:
            rows_by_id[values["record_id"]] = (values, document)
            continue

        previous_values, previous_document = previous

        if previous_values["content_hash"] == values["content_hash"]:
            continue

        previous_updated = _record_update_sort_key(previous_document["record"])
        current_updated = _record_update_sort_key(document["record"])

        if current_updated >= previous_updated:
            rows_by_id[values["record_id"]] = (values, document)

    rows = [values for values, _ in rows_by_id.values()]
    current_ids = [row["record_id"] for row in rows]

    connection_context = (
        psycopg.connect(POSTGRES_DSN) if connection is None else nullcontext(connection)
    )
    with connection_context as conn:
        with conn.cursor() as cur:
            cur.execute(
                f"""
                CREATE TABLE IF NOT EXISTS {POSTGRES_ATTENDANCE_TABLE} (
                    record_id TEXT PRIMARY KEY,
                    content_hash TEXT NOT NULL,

                    employee_id TEXT NOT NULL,
                    name TEXT NOT NULL,
                    organization_unit TEXT,
                    country TEXT,
                    work_location TEXT,
                    department TEXT,
                    position TEXT,
                    job TEXT,
                    gradeset TEXT,
                    grade TEXT,

                    attendance_date DATE NOT NULL,
                    day TEXT,
                    day_type TEXT,
                    holiday_type TEXT,
                    shift TEXT,
                    status TEXT,
                    exception TEXT,

                    total_worked_hrs DOUBLE PRECISION,
                    lateness_hrs DOUBLE PRECISION,
                    early_out_hrs DOUBLE PRECISION,
                    overbreak_hrs DOUBLE PRECISION,
                    regular_units DOUBLE PRECISION,

                    pre_ot_hrs DOUBLE PRECISION,
                    post_ot_hrs DOUBLE PRECISION,
                    total_ot DOUBLE PRECISION,
                    ot_authorized DOUBLE PRECISION,
                    ot_not_authorized DOUBLE PRECISION,

                    leave_type TEXT,
                    leave_hrs DOUBLE PRECISION,

                    last_updated_at TIMESTAMP,
                    source_file TEXT,
                    source_sheet TEXT,
                    source_excel_row INTEGER,
                    source_jsonl_line INTEGER,

                    search_text TEXT NOT NULL,
                    record_json JSONB NOT NULL,
                    raw_row_key TEXT,
                    synced_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            )
            cur.execute(
                f"ALTER TABLE {POSTGRES_ATTENDANCE_TABLE} "
                "ADD COLUMN IF NOT EXISTS raw_row_key TEXT"
            )

            # Exact-query indexes.
            cur.execute(
                f"CREATE INDEX IF NOT EXISTS idx_attendance_employee_id "
                f"ON {POSTGRES_ATTENDANCE_TABLE}(employee_id)"
            )
            cur.execute(
                f"CREATE INDEX IF NOT EXISTS idx_attendance_date "
                f"ON {POSTGRES_ATTENDANCE_TABLE}(attendance_date)"
            )
            cur.execute(
                f"CREATE INDEX IF NOT EXISTS idx_attendance_department "
                f"ON {POSTGRES_ATTENDANCE_TABLE}(department)"
            )
            cur.execute(
                f"CREATE INDEX IF NOT EXISTS idx_attendance_exception "
                f"ON {POSTGRES_ATTENDANCE_TABLE}(exception)"
            )
            cur.execute(
                f"CREATE INDEX IF NOT EXISTS idx_attendance_shift "
                f"ON {POSTGRES_ATTENDANCE_TABLE}(shift)"
            )
            upsert_sql = f"""
                INSERT INTO {POSTGRES_ATTENDANCE_TABLE} (
                    record_id,
                    content_hash,
                    employee_id,
                    name,
                    organization_unit,
                    country,
                    work_location,
                    department,
                    position,
                    job,
                    gradeset,
                    grade,
                    attendance_date,
                    day,
                    day_type,
                    holiday_type,
                    shift,
                    status,
                    exception,
                    total_worked_hrs,
                    lateness_hrs,
                    early_out_hrs,
                    overbreak_hrs,
                    regular_units,
                    pre_ot_hrs,
                    post_ot_hrs,
                    total_ot,
                    ot_authorized,
                    ot_not_authorized,
                    leave_type,
                    leave_hrs,
                    last_updated_at,
                    source_file,
                    source_sheet,
                    source_excel_row,
                    source_jsonl_line,
                    search_text,
                    record_json,
                    raw_row_key,
                    synced_at
                )
                VALUES (
                    %(record_id)s,
                    %(content_hash)s,
                    %(employee_id)s,
                    %(name)s,
                    %(organization_unit)s,
                    %(country)s,
                    %(work_location)s,
                    %(department)s,
                    %(position)s,
                    %(job)s,
                    %(gradeset)s,
                    %(grade)s,
                    %(attendance_date)s::date,
                    %(day)s,
                    %(day_type)s,
                    %(holiday_type)s,
                    %(shift)s,
                    %(status)s,
                    %(exception)s,
                    %(total_worked_hrs)s,
                    %(lateness_hrs)s,
                    %(early_out_hrs)s,
                    %(overbreak_hrs)s,
                    %(regular_units)s,
                    %(pre_ot_hrs)s,
                    %(post_ot_hrs)s,
                    %(total_ot)s,
                    %(ot_authorized)s,
                    %(ot_not_authorized)s,
                    %(leave_type)s,
                    %(leave_hrs)s,
                    %(last_updated_at)s::timestamp,
                    %(source_file)s,
                    %(source_sheet)s,
                    %(source_excel_row)s,
                    %(source_jsonl_line)s,
                    %(search_text)s,
                    %(record_json)s::jsonb,
                    %(raw_row_key)s,
                    NOW()
                )
                ON CONFLICT (record_id)
                DO UPDATE SET
                    content_hash = EXCLUDED.content_hash,
                    employee_id = EXCLUDED.employee_id,
                    name = EXCLUDED.name,
                    organization_unit = EXCLUDED.organization_unit,
                    country = EXCLUDED.country,
                    work_location = EXCLUDED.work_location,
                    department = EXCLUDED.department,
                    position = EXCLUDED.position,
                    job = EXCLUDED.job,
                    gradeset = EXCLUDED.gradeset,
                    grade = EXCLUDED.grade,
                    attendance_date = EXCLUDED.attendance_date,
                    day = EXCLUDED.day,
                    day_type = EXCLUDED.day_type,
                    holiday_type = EXCLUDED.holiday_type,
                    shift = EXCLUDED.shift,
                    status = EXCLUDED.status,
                    exception = EXCLUDED.exception,
                    total_worked_hrs = EXCLUDED.total_worked_hrs,
                    lateness_hrs = EXCLUDED.lateness_hrs,
                    early_out_hrs = EXCLUDED.early_out_hrs,
                    overbreak_hrs = EXCLUDED.overbreak_hrs,
                    regular_units = EXCLUDED.regular_units,
                    pre_ot_hrs = EXCLUDED.pre_ot_hrs,
                    post_ot_hrs = EXCLUDED.post_ot_hrs,
                    total_ot = EXCLUDED.total_ot,
                    ot_authorized = EXCLUDED.ot_authorized,
                    ot_not_authorized = EXCLUDED.ot_not_authorized,
                    leave_type = EXCLUDED.leave_type,
                    leave_hrs = EXCLUDED.leave_hrs,
                    last_updated_at = EXCLUDED.last_updated_at,
                    source_file = EXCLUDED.source_file,
                    source_sheet = EXCLUDED.source_sheet,
                    source_excel_row = EXCLUDED.source_excel_row,
                    source_jsonl_line = EXCLUDED.source_jsonl_line,
                    search_text = EXCLUDED.search_text,
                    record_json = EXCLUDED.record_json,
                    raw_row_key = EXCLUDED.raw_row_key,
                    synced_at = NOW()
                WHERE
                    {POSTGRES_ATTENDANCE_TABLE}.content_hash
                    IS DISTINCT FROM EXCLUDED.content_hash
                    OR {POSTGRES_ATTENDANCE_TABLE}.raw_row_key
                    IS DISTINCT FROM EXCLUDED.raw_row_key
            """

            for row in rows:
                params = dict(row)
                params["record_json"] = json.dumps(
                    row["record_json"],
                    ensure_ascii=False,
                    default=str,
                )
                cur.execute(upsert_sql, params)
                if cur.rowcount > 0:
                    stats.postgres_attendance_upserts += 1

            # Remove structured rows that disappeared from the current source.
            if current_ids:
                cur.execute(
                    f"DELETE FROM {POSTGRES_ATTENDANCE_TABLE} "
                    "WHERE NOT (record_id = ANY(%s))",
                    (current_ids,),
                )
            else:
                cur.execute(f"DELETE FROM {POSTGRES_ATTENDANCE_TABLE}")

        if connection is None:
            conn.commit()

    logger.info(
        "PostgreSQL attendance sync complete upserts=%s",
        stats.postgres_attendance_upserts,
    )

    return stats


def sync_chunks_to_postgres(chunks, stats=None):
    """
    Improvement 13:
    Optional PostgreSQL + pgvector mirror.

    It REUSES embeddings already stored in Chroma, so enabling PostgreSQL
    does not make a second OpenAI embedding pass.

    Enable with:
        ENABLE_POSTGRES=true
        POSTGRES_DSN=postgresql://user:password@host/db

    Requires:
        uv add "psycopg[binary]"
    and PostgreSQL pgvector extension installed.
    """
    if stats is None:
        stats = IngestionStats()

    if not ENABLE_POSTGRES:
        return stats

    if not ENABLE_PGVECTOR:
        logger.info("PostgreSQL pgvector disabled; semantic_chunks_backend=chroma")
        return stats

    if not POSTGRES_DSN:
        raise ValueError("ENABLE_POSTGRES is true but POSTGRES_DSN is empty.")

    try:
        import psycopg
    except ImportError as exc:
        raise RuntimeError(
            'PostgreSQL sync is enabled. Install psycopg with: uv add "psycopg[binary]"'
        ) from exc

    items = _prepare_embedding_items(chunks)
    part_ids = [item["id"] for item in items]

    chroma = create_chroma_client()
    collection = chroma.get_or_create_collection(COLLECTION_NAME)

    stored = collection.get(
        ids=part_ids,
        include=["documents", "metadatas", "embeddings"],
    )

    by_id = {}

    for item_id, document, metadata, embedding in zip(
        stored.get("ids", []),
        stored.get("documents", []),
        stored.get("metadatas", []),
        stored.get("embeddings", []),
    ):
        by_id[item_id] = {
            "content": document,
            "metadata": metadata or {},
            "embedding": embedding,
        }

    missing = [item_id for item_id in part_ids if item_id not in by_id]

    if missing:
        raise RuntimeError(
            "Cannot mirror to PostgreSQL because some Chroma embeddings "
            f"are missing. First missing id: {missing[0]}"
        )

    with psycopg.connect(POSTGRES_DSN) as conn:
        with conn.cursor() as cur:
            cur.execute("CREATE EXTENSION IF NOT EXISTS vector")
            cur.execute(
                f"""
                CREATE TABLE IF NOT EXISTS {POSTGRES_CHUNKS_TABLE} (
                    chunk_id TEXT PRIMARY KEY,
                    record_id TEXT NOT NULL,
                    chunk_type TEXT NOT NULL,
                    content_hash TEXT NOT NULL,
                    content TEXT NOT NULL,
                    metadata JSONB NOT NULL,
                    embedding VECTOR({PGVECTOR_DIMENSIONS}) NOT NULL,
                    updated_at TIMESTAMPTZ NOT NULL DEFAULT NOW()
                )
                """
            )
            cur.execute(
                "CREATE INDEX IF NOT EXISTS idx_knowledge_record_id "
                f"ON {POSTGRES_CHUNKS_TABLE}(record_id)"
            )
            cur.execute(
                "CREATE INDEX IF NOT EXISTS idx_knowledge_chunk_type "
                f"ON {POSTGRES_CHUNKS_TABLE}(chunk_type)"
            )

            for item in items:
                stored_item = by_id[item["id"]]

                vector_literal = (
                    "["
                    + ",".join(str(value) for value in stored_item["embedding"])
                    + "]"
                )

                cur.execute(
                    f"""
                    INSERT INTO {POSTGRES_CHUNKS_TABLE} (
                        chunk_id,
                        record_id,
                        chunk_type,
                        content_hash,
                        content,
                        metadata,
                        embedding,
                        updated_at
                    )
                    VALUES (%s, %s, %s, %s, %s, %s::jsonb, %s::vector, NOW())
                    ON CONFLICT (chunk_id)
                    DO UPDATE SET
                        record_id = EXCLUDED.record_id,
                        chunk_type = EXCLUDED.chunk_type,
                        content_hash = EXCLUDED.content_hash,
                        content = EXCLUDED.content,
                        metadata = EXCLUDED.metadata,
                        embedding = EXCLUDED.embedding,
                        updated_at = NOW()
                    """,
                    (
                        item["id"],
                        item["record_id"],
                        item["chunk_type"],
                        item["content_hash"],
                        stored_item["content"],
                        json.dumps(
                            stored_item["metadata"],
                            ensure_ascii=False,
                        ),
                        vector_literal,
                    ),
                )

                stats.postgres_upserts += 1

            cur.execute(
                f"DELETE FROM {POSTGRES_CHUNKS_TABLE} WHERE NOT (chunk_id = ANY(%s))",
                (part_ids,),
            )

        conn.commit()

    return stats


def _count_invalid_records():
    if not INVALID_JSONL_PATH.exists():
        return 0

    with open(INVALID_JSONL_PATH, "r", encoding="utf-8") as handle:
        return sum(1 for line in handle if line.strip())


def _print_stats(stats):
    logger.info(
        "Ingestion summary jsonl_records=%s invalid_records=%s "
        "attendance_chunks=%s period_chunks=%s new_records=%s "
        "changed_records=%s unchanged_records=%s deleted_records=%s "
        "duplicate_records=%s embedding_inputs=%s embedding_api_batches=%s "
        "chroma_upserts=%s metadata_updates=%s postgres_chunks=%s "
        "postgres_attendance=%s raw_rows_seen=%s raw_rows_inserted=%s "
        "unknown_partitions=%s quarantined_rows=%s "
        "stage_seconds=%s elapsed_seconds=%.2f",
        stats.jsonl_records,
        stats.invalid_records,
        stats.attendance_chunks,
        stats.period_chunks,
        stats.new_records,
        stats.changed_records,
        stats.unchanged_records,
        stats.deleted_records,
        stats.duplicate_records,
        stats.embedding_inputs,
        stats.embedding_api_batches,
        stats.chroma_upserts,
        stats.metadata_updates,
        stats.postgres_upserts,
        stats.postgres_attendance_upserts,
        stats.raw_rows_seen,
        stats.raw_rows_inserted,
        stats.unknown_partitions,
        stats.quarantined_rows,
        stats.stage_seconds,
        stats.elapsed_seconds,
    )


def _journal_documents(documents, ledger=None):
    ledger = ledger or IngestionLedger(settings.ingestion_state_path)
    source_hash = _content_hash(_source_signature(_source_files()))
    with ledger.writer() as writer:
        generation = writer.begin_or_resume_generation(source_hash)
        for document in documents:
            record = document["record"]
            business = _business_record(record)
            search_text = _record_to_search_text(record)
            metadata = _record_to_metadata(document)
            writer.upsert_record(
                generation,
                _make_record_id(record),
                business,
                business_hash=_content_hash(business),
                embedding_hash=_content_hash(
                    {
                        "model": EMBEDDING_MODEL,
                        "index_schema_version": settings.index_schema_version,
                        "text": search_text,
                    }
                ),
                metadata_hash=_content_hash(metadata),
                source_locator=(
                    f"{document.get('source')}#{document.get('sheet')}!"
                    f"{document.get('excel_row')}"
                ),
            )
    return ledger, generation


def _run_ingestion(ledger):
    started = perf_counter()
    stats = IngestionStats()

    stage_started = perf_counter()
    snapshot = load_source_snapshot(ledger)
    documents = list(snapshot.documents)
    stats.stage_seconds["jsonl_read_validation"] = perf_counter() - stage_started
    stats.jsonl_records = len(documents)
    stats.invalid_records = snapshot.invalid_count
    stats.unknown_partitions = sum(
        1 for item in snapshot.partition_summaries if item.get("domain") == "unknown"
    )
    stats.quarantined_rows = sum(
        1 for row in snapshot.raw_rows if row.status == "quarantined"
    )

    if stats.unknown_partitions and not ENABLE_POSTGRES:
        raise RuntimeError(
            "Unknown source partitions were quarantined, but PostgreSQL is disabled; "
            "refusing Chroma mutation because the raw rows cannot be preserved."
        )
    if snapshot.unsafe_reasons:
        if snapshot.raw_rows and ENABLE_POSTGRES:
            stats.raw_rows_seen = len(snapshot.raw_rows)
            stats.raw_rows_inserted = store_raw_rows_to_postgres(snapshot.raw_rows)
        raise RuntimeError(
            "Unsafe attendance source snapshot; typed attendance and Chroma were not "
            "changed: " + "; ".join(snapshot.unsafe_reasons)
        )
    if not documents:
        if snapshot.raw_rows and ENABLE_POSTGRES:
            stats.raw_rows_seen = len(snapshot.raw_rows)
            stats.raw_rows_inserted = store_raw_rows_to_postgres(snapshot.raw_rows)
        raise RuntimeError(
            "Ingestion produced zero valid attendance records; refusing destructive "
            "sink synchronization."
        )
    if stats.invalid_records:
        if snapshot.raw_rows and ENABLE_POSTGRES:
            stats.raw_rows_seen = len(snapshot.raw_rows)
            stats.raw_rows_inserted = store_raw_rows_to_postgres(snapshot.raw_rows)
        raise RuntimeError(
            f"Ingestion found {stats.invalid_records} invalid attendance record(s); "
            "refusing sink synchronization because a partial snapshot could delete "
            "previously valid records."
        )

    documents = canonicalize_documents(documents, stats)
    stats.jsonl_records = len(documents)

    stage_started = perf_counter()
    attendance_chunks = create_record_chunks(documents)
    stats.stage_seconds["attendance_chunking"] = perf_counter() - stage_started
    stats.attendance_chunks = len(attendance_chunks)

    period_chunks = []
    if ENABLE_EMPLOYEE_PERIOD_CHUNKS:
        stage_started = perf_counter()
        period_chunks = create_employee_period_chunks(documents)
        stats.stage_seconds["period_chunking"] = perf_counter() - stage_started
        stats.period_chunks = len(period_chunks)

    all_chunks = attendance_chunks + period_chunks

    generation = None
    if documents:
        ledger, generation = _journal_documents(documents, ledger)

    stage_started = perf_counter()
    publish_postgres_snapshot(
        documents,
        snapshot.raw_rows,
        generation or "",
        stats,
    )
    stats.stage_seconds["postgres_attendance_sync"] = perf_counter() - stage_started

    stage_started = perf_counter()
    sync_embeddings_to_chroma(
        all_chunks,
        stats,
        ledger=ledger,
        generation=generation,
    )
    stats.stage_seconds["embedding_and_chroma_sync"] = perf_counter() - stage_started

    stage_started = perf_counter()
    sync_chunks_to_postgres(all_chunks, stats)
    stats.stage_seconds["postgres_vector_sync"] = perf_counter() - stage_started

    if generation is not None:
        with ledger.writer() as writer:
            writer.activate_generation(generation)

    stats.elapsed_seconds = perf_counter() - started
    _print_stats(stats)

    logger.info("Ingestion complete")


def main():
    logging.basicConfig(
        level=getattr(logging, settings.log_level, logging.INFO),
        format="%(asctime)s %(levelname)s %(name)s %(message)s",
    )
    ledger = IngestionLedger(settings.ingestion_state_path)
    with ledger.run_lock():
        return _run_ingestion(ledger)


if __name__ == "__main__":
    main()
